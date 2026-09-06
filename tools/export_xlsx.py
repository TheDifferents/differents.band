#!/usr/bin/env python3
"""Export the song book + wishlist to Excel.

    python3 tools/export_xlsx.py            -> differents-songbook.xlsx

Sheet 1 "Dance Numbers" is the working sheet: every dance-floor song we play
plus every song on the wishlist, one list, slowest to fastest.
Sheet 2 "Full Book" is everything, same sort, for reference.

The dance tiering below is a judgment call, not data — edit these two lists to
change what lands on sheet 1.
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# The band's own call, from what has actually worked on a floor — not a guess.
# 18 songs were cut from an earlier draft of these lists because they have never
# got a floor going live, whatever their tempo or reputation suggests.
TIER1 = {"I Will Survive", "Superstition", "Billie Jean", "9 to 5",
         "Walking On Sunshine", "Burning Love", "What I Like About You",
         "Mr. Brightside"}
TIER2 = {"Crazy Little Thing Called Love", "Valerie", "Don't You Forget About Me",
         "Low Rider", "Jumpin' Jack Flash", "Just What I Needed", "Time Warp"}
# Cut as unproven on a floor: Use Me, The Seeker, Dani California, Sweet Emotion,
# Mysterious Ways, Hard To Handle, Spooky, Mustang Sally, Moneygrabber, Get Lucky,
# Chain Of Fools, Some Kind of Wonderful, And She Was, Two Tickets To Paradise,
# Pretty Vegas, Molly's Chambers, Lonely Boy, Rock And Roll.

def load(path):
    out = []
    for line in open(path, encoding='utf-8'):
        line = line.strip()
        if line and not line.startswith('#'):
            out.append(json.loads(line))
    return out

book = [s for s in load('data/songs.jsonl') if s['status'] in ('rotation', 'bullpen')]
wish = load('data/wishlist.jsonl')

rows = []
for s in book:
    rows.append({**s, 'list': 'Book',
                 'dance': 'Fills the floor' if s['title'] in TIER1
                          else 'Holds the floor' if s['title'] in TIER2 else ''})
for s in wish:
    rows.append({**s, 'chords': None, 'duration': None, 'sets': [],
                 'status': 'wishlist', 'list': 'Wishlist', 'dance': 'Proposed'})
rows.sort(key=lambda s: (s['bpm'], s['title']))
dance = [r for r in rows if r['dance']]

COLS = [('BPM', 'bpm', 7), ('Title', 'title', 30), ('Artist', 'artist', 26),
        ('Year', 'year', 7), ('List', 'list', 10), ('Status', 'status', 11),
        ('Dance floor', 'dance', 16), ('Sets', 'sets', 17),
        ('Chords', 'chords', 22), ('Length', 'duration', 8), ('Fact', 'fact', 82)]

INK   = '14100E'
GOLD  = 'D4A91F'
CREAM = 'F4EFE6'
WISHF = PatternFill('solid', fgColor='FBF3D8')
HEADF = PatternFill('solid', fgColor=INK)
thin  = Side(style='thin', color='DDD5C6')

def sheet(wb, title, data, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append([c[0] for c in COLS])
    for i, (label, _, width) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color=CREAM, size=11)
        c.fill = HEADF
        c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    for r in data:
        ws.append([', '.join(r[k]) if k == 'sets' else r[k] for _, k, _ in COLS])
        n = ws.max_row
        wishy = r['list'] == 'Wishlist'
        for i, (_, key, _) in enumerate(COLS, 1):
            c = ws.cell(row=n, column=i)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical='top',
                                    wrap_text=(key == 'fact'),
                                    horizontal='center' if key in ('bpm', 'year') else 'left')
            if wishy:
                c.fill = WISHF
            if key == 'bpm':
                c.font = Font(bold=True, color='8A6A08' if wishy else '3A3A3A')
            if key == 'title':
                c.font = Font(bold=True)
        ws.row_dimensions[n].height = 30

    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{ws.max_row}'
    ws.sheet_view.showGridLines = False
    return ws

wb = Workbook()
sheet(wb, 'Dance Numbers', dance, first=True)
sheet(wb, 'Full Book', rows)
out = 'differents-songbook.xlsx'
wb.save(out)
print(f'{out}')
print(f'  Dance Numbers  {len(dance):3d} rows  '
      f'({sum(1 for r in dance if r["list"]=="Book")} ours, '
      f'{sum(1 for r in dance if r["list"]=="Wishlist")} proposed)  '
      f'{dance[0]["bpm"]}–{dance[-1]["bpm"]} BPM')
print(f'  Full Book      {len(rows):3d} rows  {rows[0]["bpm"]}–{rows[-1]["bpm"]} BPM')
